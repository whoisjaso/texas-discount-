#!/usr/bin/env python3
"""Build the bundled EPA equivalent-test-weight table the desk estimates empty weight from.

What it is
  EPA publishes, every model year, the "Test Car List": every vehicle a
  manufacturer tested for fuel economy, with its Equivalent Test Weight (ETW).
  ETW is the inertia class of the loaded vehicle weight, and loaded vehicle
  weight is curb weight plus 300 lb (40 CFR 86.1803-01; 40 CFR 1066.805,
  Table 1 and paragraph (b)). So "ETW less 300 lb" is an estimate of curb
  weight, within half a class (62.5, 125 or 250 lb).

  This script reads every Test Car List file for MY1995 to MY2026, groups the
  test rows by (year, make or manufacturer group, model, displacement, drive,
  hybrid/EV), keeps EVERY ETW of each group (with repeats, so the median the
  desk takes matches reference_estimate.py exactly) and writes:

    src/lib/vehicles/empty-weight/epa-etw-table.generated.js

  as `export default {...}`, with the source URL, sha256 and download date of
  every file it read, so the desk can say where an estimate came from.

Usage
  python3 build_epa_table.py --download [--cache /tmp/epa-test-car]
      fetch the EPA page, collect every test car file link, download them to
      the cache (outside the repo), record each file's URL, sha256 and the
      date it was downloaded in CACHE/manifest.json, and build from there.
  python3 build_epa_table.py --from-dir DIR
      build from files already downloaded, reading DIR/manifest.json.
  python3 build_epa_table.py --from-dir DIR --urls DIR/urls.txt --downloaded YYYY-MM-DD
      a folder with no manifest yet: write one from urls.txt (one URL per
      line) and the date the files were fetched, then build.

Reproducible: the dates come from the manifest, never from file times, and
"built" is the newest download date in it, so the same files and manifest
always give the same table byte for byte. A file the manifest does not name,
or whose sha256 no longer matches it, stops the build: every row of the table
must say exactly which EPA file it came from.

The yearly step
  When EPA posts a new model year, run `--download`, commit the regenerated
  table and fixture (make_fixtures.py), and run the desk's tests.

Requires: Python 3.9+, openpyxl (pip install openpyxl). No npm install.
"""
import argparse
import collections
import csv
import datetime
import glob
import hashlib
import io
import json
import os
import re
import sys
import urllib.request
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
DESK = os.path.normpath(os.path.join(HERE, '..', '..'))
OUT = os.path.join(DESK, 'src', 'lib', 'vehicles', 'empty-weight', 'epa-etw-table.generated.js')
PAGE = 'https://www.epa.gov/compliance-and-fuel-economy-data/data-cars-used-testing-fuel-economy'
FIRST_YEAR, LAST_YEAR = 1995, 2026

# The same constants the lookup uses (reference_estimate.py imports them from here).
FOUR = {'4WD', '4X4', 'AWD', '4MATIC', 'XDRIVE', 'QUATTRO', '4MOTION', 'SYNCRO', 'SHAWD', 'ALL4', '4XE', 'E4WD', 'AWD/4WD'}
TWO = {'2WD', '4X2', 'FWD', 'RWD'}
# make -> regex over EPA manufacturer names (the 2000-2009 files have no make column)
GROUP = [
    (r'GENERAL MOTORS|^GM|GMOLS|DAEWOO|QUANTUM', {'CHEVROLET', 'GMC', 'BUICK', 'CADILLAC', 'PONTIAC', 'OLDSMOBILE', 'SATURN', 'HUMMER', 'GEO', 'DAEWOO', 'SAAB'}),
    (r'FORD|ROUSH|SALEEN|SHELBY', {'FORD', 'LINCOLN', 'MERCURY', 'MAZDA'}),
    (r'CHRYS|DAIMLERCHRYSLER|FCA|CHRYSLER', {'CHRYSLER', 'DODGE', 'JEEP', 'PLYMOUTH', 'EAGLE', 'RAM'}),
    (r'TOYOT|NEW UNITED|NUMMI', {'TOYOTA', 'LEXUS', 'SCION', 'PONTIAC', 'CHEVROLET', 'GEO'}),
    (r'HONDA', {'HONDA', 'ACURA'}), (r'NISS', {'NISSAN', 'INFINITI'}), (r'HYND|HYUNDAI', {'HYUNDAI', 'GENESIS'}),
    (r'KIA', {'KIA'}), (r'MAZDA|MZM', {'MAZDA'}), (r'MITS|MMMA|MMAL', {'MITSUBISHI', 'CHRYSLER', 'DODGE', 'EAGLE', 'PLYMOUTH'}),
    (r'FUJI', {'SUBARU'}), (r'V W|VOLKSWAGEN|AUDI', {'VOLKSWAGEN', 'AUDI'}), (r'BMW', {'BMW', 'MINI'}),
    (r'MBZ|MERCEDES', {'MERCEDES-BENZ', 'SMART'}), (r'VOLVO', {'VOLVO'}), (r'SAAB', {'SAAB'}), (r'SUZUK', {'SUZUKI', 'CHEVROLET', 'GEO'}),
    (r'ISUZU', {'ISUZU', 'HONDA'}), (r'PRSCH|PORSCHE', {'PORSCHE'}), (r'JAGUAR', {'JAGUAR'}), (r'LAND ROVER|LAROV|ROVGP', {'LAND ROVER'}),
]
MAKE_ALIAS = {'MERCEDES BENZ': 'MERCEDES-BENZ', 'MECEDES-BENZ': 'MERCEDES-BENZ', 'VW': 'VOLKSWAGEN', 'ROLLS ROYCE': 'ROLLS-ROYCE',
              'LINCOLN-MERCURY': 'LINCOLN', 'RANGE ROVER': 'LAND ROVER', 'SRT': 'DODGE'}
HYBRID_RX = re.compile(r'HYBRID|HEV|PHEV|PLUG-IN|4XE|ENERGI|E-TRON')
DASHES = re.compile('[‒–—―−]')


def toks(s):
    return [t for t in re.split(r'[^A-Z0-9]+', (s or '').upper()) if t]


def alnum(s):
    return re.sub(r'[^A-Z0-9]', '', (s or '').upper())


def drive2or4(code, text):
    t = set(toks(text)) | {alnum(x) for x in (text or '').upper().split()}
    if t & {alnum(x) for x in FOUR}:
        return 4
    if code in ('4', 'A', 'P'):
        return 4
    if t & TWO or code in ('F', 'R'):
        return 2
    return None


def f2(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def i2(x):
    try:
        return int(float(x))
    except (TypeError, ValueError):
        return None


def drv_norm(code, desc=''):
    c = (code or '').strip().upper()
    d = (desc or '').lower()
    if c in ('F', 'FWD'):
        return 'F'
    if c in ('R', 'RWD'):
        return 'R'
    if c in ('4', '4WD', 'A', 'AWD', 'P', 'PT4', '4WD/AWD'):
        return c if c in ('A', 'P') else '4'
    if 'front' in d:
        return 'F'
    if 'rear' in d:
        return 'R'
    if 'all' in d:
        return 'A'
    if '4' in d or 'four' in d:
        return '4'
    return c or ''


# ---------------------------------------------------------------- reading the files, by era

def read_fixed_width(name, data):
    """MY1984-1997: fixed-width text inside a zip."""
    rows = []
    for line in data.decode('latin1').splitlines():
        if len(line) < 200 or not line[:4].isdigit():
            continue
        cid = i2(line[14:18])
        rows.append(dict(year=int(line[0:4]), mfr=line[83:88].strip(), make='', model=line[165:190].strip(),
                         disp_l=round(cid * 0.016387, 1) if cid else None, drive=drv_norm(line[203:206]),
                         etw=i2(line[50:54]), src=name))
    return rows


def read_csv_9899(name, data):
    rows = []
    for r in csv.DictReader(io.StringIO(data.decode('latin1'))):
        cid = i2(r['cid'])
        rows.append(dict(year=i2(r['yr']), mfr=r['mfr name'].strip(), make='', model=r['carline'].strip(),
                         disp_l=round(cid * 0.016387, 1) if cid else None, drive=drv_norm(r['drv']),
                         etw=i2(r['etw']), src=name))
    return rows


def modern(r, src):
    def g(k):
        return r.get(k) if r.get(k) is not None else ''
    return dict(year=i2(g('Model Year')), mfr=str(g('Vehicle Manufacturer Name')).strip(),
                make=str(g('Represented Test Veh Make')).strip(), model=str(g('Represented Test Veh Model')).strip(),
                disp_l=f2(g('Test Veh Displacement (L)')),
                drive=drv_norm(str(g('Drive System Code')), str(g('Drive System Description'))),
                etw=i2(g('Equivalent Test Weight (lbs.)')), src=src)


def read_csv_modern(name, data):
    rd = csv.DictReader(io.StringIO(data.decode('latin1')))
    rows = []
    if 'MDLYR_DT' in (rd.fieldnames or []):  # 2000-2009: no make column
        for r in rd:
            cid = i2(r['GBE_CID_MSR'])
            rows.append(dict(year=i2(r['MDLYR_DT']), mfr=r['VI_MFR_NM'].strip(), make='', model=r['CL_NM'].strip(),
                             disp_l=round(cid * 0.016387, 1) if cid else None, drive=drv_norm(r['DRV_SYS_CD']),
                             etw=i2(r['VC_DSN_ETW_MSR']), src=name))
    else:
        for r in rd:
            rows.append(modern(r, name))
    return rows


def read_xlsx(name, path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    hdr = [str(h).strip() if h else '' for h in next(it)]
    rows = []
    for v in it:
        if v[0] is None:
            continue
        rows.append(modern(dict(zip(hdr, v)), name))
    return rows


def year_of(name):
    m = re.match(r'(\d\d)', os.path.basename(name))
    if not m:
        return None
    yy = int(m.group(1))
    return 1900 + yy if yy >= 84 else 2000 + yy


def read_all(files):
    rows = []
    for path in files:
        name = os.path.basename(path)
        y = year_of(name)
        if y is None or y < FIRST_YEAR or y > LAST_YEAR:
            continue
        if name.lower().endswith('.zip'):
            z = zipfile.ZipFile(path)
            for inner in z.namelist():
                data = z.read(inner)
                if inner.lower().endswith('.csv'):
                    rows += read_csv_9899(name, data)
                else:
                    rows += read_fixed_width(name, data)
        elif name.lower().endswith('.csv'):
            rows += read_csv_modern(name, open(path, 'rb').read())
        elif name.lower().endswith('.xlsx'):
            rows += read_xlsx(name, path)
    return rows


# ---------------------------------------------------------------- downloading

def fetch(url):
    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0 (desk empty-weight table build)'})
    with urllib.request.urlopen(req, timeout=120) as res:
        return res.read()


MANIFEST = 'manifest.json'


def sha256_of(path):
    return hashlib.sha256(open(path, 'rb').read()).hexdigest()


def read_manifest(src_dir):
    path = os.path.join(src_dir, MANIFEST)
    if not os.path.exists(path):
        return None
    doc = json.load(open(path))
    return {e['file']: e for e in doc['files']}


def write_manifest(src_dir, entries):
    doc = dict(page=PAGE, files=[entries[k] for k in sorted(entries)])
    open(os.path.join(src_dir, MANIFEST), 'w').write(json.dumps(doc, indent=1) + '\n')


def download(cache):
    os.makedirs(cache, exist_ok=True)
    html = fetch(PAGE).decode('utf8', 'replace')
    links = sorted(set(re.findall(r'href="([^"]+?(?:tstcar|testcar|mftcl)[^"]*?\.(?:csv|xlsx|zip))"', html, re.I)))
    urls = [l if l.startswith('http') else 'https://www.epa.gov' + l for l in links]
    known = read_manifest(cache) or {}
    today = datetime.date.today().isoformat()
    entries = {}
    for u in urls:
        name = os.path.basename(u)
        dest = os.path.join(cache, name)
        if not os.path.exists(dest):
            print('downloading', u)
            open(dest, 'wb').write(fetch(u))
        digest = sha256_of(dest)
        old = known.get(name)
        # The same bytes keep the date they were first fetched; new bytes get today.
        stamp = old['downloaded'] if old and old.get('sha256') == digest and old.get('url') == u else today
        entries[name] = dict(file=name, url=u, sha256=digest, downloaded=stamp)
    write_manifest(cache, entries)
    open(os.path.join(cache, 'urls.txt'), 'w').write('\n'.join(urls) + '\n')
    return cache


def manifest_from_urls(src_dir, urls_file, downloaded):
    """A first manifest for a folder fetched by hand: URL per file, one stated date."""
    if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', downloaded or ''):
        sys.exit('--downloaded YYYY-MM-DD is required with --urls: the date the files were fetched')
    url_by_name = {}
    for line in open(urls_file):
        u = line.strip()
        if u:
            url_by_name[os.path.basename(u)] = u
    entries = {}
    for path in table_files(src_dir):
        name = os.path.basename(path)
        if name not in url_by_name:
            sys.exit('no URL for %s in %s: every file must say where it came from' % (name, urls_file))
        entries[name] = dict(file=name, url=url_by_name[name], sha256=sha256_of(path), downloaded=downloaded)
    write_manifest(src_dir, entries)
    return entries


def table_files(src_dir):
    out = []
    for path in sorted(glob.glob(os.path.join(src_dir, '*'))):
        name = os.path.basename(path)
        y = year_of(name)
        if y is None or y < FIRST_YEAR or y > LAST_YEAR or not re.search(r'\.(csv|xlsx|zip)$', name, re.I):
            continue
        out.append(path)
    return out


# ---------------------------------------------------------------- the table

def group_key(mfr):
    """A make-less row matches every make of every group its manufacturer name matches."""
    hits = [i for i, (rx, _) in enumerate(GROUP) if re.search(rx, mfr, re.I)]
    if not hits:
        return None, None
    key = '@' + '+'.join(GROUP[i][0].split('|')[0].replace('^', '') for i in hits)
    makes = sorted(set().union(*(GROUP[i][1] for i in hits)))
    return key, makes


def clean_model(model):
    # No en or em dash ever reaches the desk's copy (no-banned-dashes test).
    return DASHES.sub('-', (model or '').strip().upper())


def build(src_dir, manifest, built):
    files = sorted(glob.glob(os.path.join(src_dir, '*')))
    sources = collections.defaultdict(list)
    for path in table_files(src_dir):
        name = os.path.basename(path)
        entry = manifest.get(name)
        if entry is None:
            sys.exit('%s is not in %s: every file must say where it came from' % (name, MANIFEST))
        digest = sha256_of(path)
        if digest != entry['sha256']:
            sys.exit('%s changed since it was recorded (sha256 %s, manifest %s): download it again' % (name, digest, entry['sha256']))
        sources[str(year_of(name))].append(dict(url=entry['url'], file=name, sha256=digest, downloaded=entry['downloaded']))

    agg = collections.defaultdict(list)
    groups = {}
    dropped = 0
    for r in read_all(files):
        y, etw = r['year'], r['etw']
        if y is None or etw is None or not (FIRST_YEAR <= y <= LAST_YEAR) or not (1000 <= etw <= 10000):
            continue
        make = MAKE_ALIAS.get(r['make'].upper().strip(), r['make'].upper().strip())
        if not make:
            key, makes = group_key(r['mfr'])
            if key is None:
                dropped += 1  # no make and no group: no vPIC make can ever match it
                continue
            groups[key] = makes
            make = key
        d = r['disp_l']
        d = round(float(d), 1) if d is not None else None
        ev = d is None or d < 0.1 or d > 50
        model = clean_model(r['model'])
        hyb = bool(HYBRID_RX.search(model))
        drive = drive2or4(r['drive'], r['model']) or 0
        flags = (1 if hyb else 0) | (2 if ev else 0)
        agg[(y, make, model, None if ev else d, drive, flags)].append(etw)

    years = collections.defaultdict(list)
    for (y, make, model, d, drive, flags), etws in sorted(agg.items(), key=lambda kv: (kv[0][0], kv[0][1], kv[0][2], str(kv[0][3]), kv[0][4], kv[0][5])):
        years[str(y)].append([make, model, d, drive, flags, sorted(etws)])

    table = dict(
        v=1,
        built=built,
        relation='curb weight estimate = ETW less 300 lb; ETW is the inertia class of loaded vehicle weight = curb + 300 lb (40 CFR 86.1803-01, 1066.805)',
        rowFormat=['make or @manufacturer group', 'EPA represented model (upper case)', 'displacement L (null for EV)',
                   'drive 0 unknown, 2, 4', 'flags 1 hybrid, 2 EV', 'every ETW tested in the group, sorted, with repeats'],
        page=PAGE,
        sources={k: sources[k] for k in sorted(sources)},
        groups={k: groups[k] for k in sorted(groups)},
        years={k: years[k] for k in sorted(years)},
    )
    return table, dropped


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('--download', action='store_true')
    ap.add_argument('--cache', default='/tmp/epa-test-car')
    ap.add_argument('--from-dir')
    ap.add_argument('--urls', help='with --downloaded, write a first manifest for a folder that has none')
    ap.add_argument('--downloaded', help='YYYY-MM-DD the files in --from-dir were fetched (with --urls)')
    ap.add_argument('--built', default=None, help='build date to record (default: newest download date in the manifest)')
    ap.add_argument('--out', default=OUT)
    a = ap.parse_args()
    if a.download:
        src = download(a.cache)
    elif a.from_dir:
        src = a.from_dir
    else:
        ap.error('give --download or --from-dir DIR')
    manifest = read_manifest(src)
    if manifest is None:
        if not a.urls:
            sys.exit('%s has no %s: give --urls FILE --downloaded YYYY-MM-DD to write one' % (src, MANIFEST))
        manifest = manifest_from_urls(src, a.urls, a.downloaded)
    used = [manifest[os.path.basename(p)]['downloaded'] for p in table_files(src) if os.path.basename(p) in manifest]
    built = a.built or (max(used) if used else None)
    if not built:
        sys.exit('no source files in %s' % src)
    table, dropped = build(src, manifest, built)
    body = json.dumps(table, separators=(',', ':'), ensure_ascii=True)
    header = ('// GENERATED by scripts/empty-weight/build_epa_table.py from the EPA Test Car List. Do not edit by hand.\n'
              '// Rebuild yearly: python3 scripts/empty-weight/build_epa_table.py --download\n'
              '// One JSON string, parsed once: faster to load than an object literal this size, and one\n'
              '// token to any tool that reads the source (the repository guards parse every file).\n')
    os.makedirs(os.path.dirname(a.out), exist_ok=True)
    open(a.out, 'w').write(header + 'export default JSON.parse(' + json.dumps(body) + ');\n')
    n = sum(len(v) for v in table['years'].values())
    print('groups', n, 'bytes', len(body), 'dropped (no make, no group)', dropped)
    print('per year', {y: len(v) for y, v in table['years'].items()})
    missing = [str(y) for y in range(FIRST_YEAR, LAST_YEAR + 1) if not table['years'].get(str(y)) or not table['sources'].get(str(y))]
    if missing:
        print('MISSING YEARS', missing)
        sys.exit(1)


if __name__ == '__main__':
    main()
